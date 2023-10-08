#%%
from docx import Document as Doc
from docx.shared import Inches

class Document():
    def __init__(self):
        self = Doc()

    def add_dataframe(self, df):
        t = self.add_table(df.shape[0]+1, df.shape[1])
        for j in range(df.shape[-1]):
            t.cell(0,j).text = df.columns[j]
        for i in range(df.shape[0]):
            for j in range(df.shape[-1]):
                t.cell(i+1,j).text = str(df.values[i,j])

        return t         
    
#%%
document = Document()
# %%
